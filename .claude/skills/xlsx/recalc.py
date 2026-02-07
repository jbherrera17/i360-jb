#!/usr/bin/env python3
"""
Recalculate Excel formulas using LibreOffice and check for errors.

Usage:
    python recalc.py <excel_file> [timeout_seconds]

Requirements:
    - LibreOffice installed
    - openpyxl: pip install openpyxl
"""

import json
import os
import platform
import subprocess
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print(json.dumps({"error": "openpyxl not installed. Run: pip install openpyxl"}))
    sys.exit(1)


EXCEL_ERRORS = ['#VALUE!', '#DIV/0!', '#REF!', '#NAME?', '#NULL!', '#NUM!', '#N/A']


def get_macro_path():
    """Get the LibreOffice macro directory path."""
    if platform.system() == 'Darwin':
        return os.path.expanduser('~/Library/Application Support/LibreOffice/4/user/basic/Standard')
    return os.path.expanduser('~/.config/libreoffice/4/user/basic/Standard')


def setup_macro():
    """Create LibreOffice macro for recalculation if needed."""
    macro_dir = get_macro_path()
    macro_file = os.path.join(macro_dir, 'Module1.xba')

    # Check if macro already exists
    if os.path.exists(macro_file):
        with open(macro_file, 'r') as f:
            if 'RecalculateAndSave' in f.read():
                return True

    # Initialize LibreOffice to create config directories
    if not os.path.exists(macro_dir):
        subprocess.run(['soffice', '--headless', '--terminate_after_init'],
                      capture_output=True, timeout=10)
        os.makedirs(macro_dir, exist_ok=True)

    # Write the macro
    macro_xml = '''<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>'''

    try:
        with open(macro_file, 'w') as f:
            f.write(macro_xml)
        return True
    except Exception:
        return False


def run_recalculation(filepath, timeout):
    """Run LibreOffice to recalculate formulas."""
    abs_path = str(Path(filepath).absolute())

    cmd = [
        'soffice', '--headless', '--norestore',
        'vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application',
        abs_path
    ]

    # Add timeout command if available
    if platform.system() == 'Darwin':
        try:
            subprocess.run(['gtimeout', '--version'], capture_output=True, timeout=1)
            cmd = ['gtimeout', str(timeout)] + cmd
        except (FileNotFoundError, subprocess.TimeoutExpired):
            pass
    elif platform.system() == 'Linux':
        cmd = ['timeout', str(timeout)] + cmd

    result = subprocess.run(cmd, capture_output=True, text=True)
    return result.returncode == 0 or result.returncode == 124


def scan_for_errors(filepath):
    """Scan all cells for Excel formula errors."""
    wb = load_workbook(filepath, data_only=True)

    errors = {err: [] for err in EXCEL_ERRORS}
    total = 0

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    for err in EXCEL_ERRORS:
                        if err in cell.value:
                            errors[err].append(f"{sheet_name}!{cell.coordinate}")
                            total += 1
                            break

    wb.close()
    return errors, total


def count_formulas(filepath):
    """Count total formulas in the workbook."""
    wb = load_workbook(filepath, data_only=False)
    count = 0

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
                    count += 1

    wb.close()
    return count


def recalc(filepath, timeout=30):
    """
    Recalculate formulas and check for errors.

    Returns dict with status, error counts, and locations.
    """
    path = Path(filepath)
    if not path.exists():
        return {'error': f'File not found: {filepath}'}

    if not setup_macro():
        return {'error': 'Failed to setup LibreOffice macro'}

    if not run_recalculation(filepath, timeout):
        return {'error': 'Recalculation failed'}

    try:
        errors, total_errors = scan_for_errors(filepath)
        formula_count = count_formulas(filepath)

        result = {
            'status': 'success' if total_errors == 0 else 'errors_found',
            'total_errors': total_errors,
            'total_formulas': formula_count,
        }

        if total_errors > 0:
            result['error_summary'] = {}
            for err_type, locations in errors.items():
                if locations:
                    result['error_summary'][err_type] = {
                        'count': len(locations),
                        'locations': locations[:20]
                    }

        return result

    except Exception as e:
        return {'error': str(e)}


def main():
    if len(sys.argv) < 2:
        print("Usage: python recalc.py <excel_file> [timeout_seconds]")
        print("\nRecalculates formulas using LibreOffice and checks for errors.")
        print("\nReturns JSON with:")
        print("  - status: 'success' or 'errors_found'")
        print("  - total_errors: Count of formula errors")
        print("  - total_formulas: Number of formulas in file")
        print("  - error_summary: Breakdown by error type with cell locations")
        sys.exit(1)

    filepath = sys.argv[1]
    timeout = int(sys.argv[2]) if len(sys.argv) > 2 else 30

    result = recalc(filepath, timeout)
    print(json.dumps(result, indent=2))


if __name__ == '__main__':
    main()
